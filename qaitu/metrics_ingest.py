"""Conservative report ingestion: explicit periods/owners, source-backed numbers."""
from __future__ import annotations

import calendar
import hashlib
import io
import math
import re
import zipfile
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

from .metrics_catalog import METRIC_CATALOG, normalize_label, resolve_metric
from .metrics_models import IngestionResult, MetricAmbiguity, MetricDefinition, MetricSource, MetricValue

MAX_FILE_BYTES = 25 * 1024 * 1024
_MONTHS = {"янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "мая": 5, "июн": 6, "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12}
_SCOPE_ALIASES = {"DNM": "ДНМ", "DKKM": "ДККМ", "DITAAD": "ДИТААД", "DOA": "ДОА", "BVA": "BVA", "БВА": "BVA"}


def parse_period(value: object) -> tuple[str, str, str] | None:
    if isinstance(value, (date, datetime)):
        if value.day != 1:
            return None
        value = value.strftime("%Y-%m")
    text = str(value or "").strip().lower().replace("ё", "е")
    if re.search(r"утвержд|протокол|приказ|редакци", text):
        return None
    if re.search(r"(?<!\d)\d{4}-\d{2}-\d{2}(?!\d)|(?<!\d)\d{1,2}[./]\d{1,2}[./]\d{4}(?!\d)", text):
        return None
    if len(set(re.findall(r"(?<!\d)((?:19|20)\d{2})(?!\d)", text))) > 1:
        return None
    quarter = re.search(r"(?:q\s*([1-4])\s*[, ._/-]*((?:19|20)\d{2})|([1-4])\s*(?:кв(?:артал)?\.?|quarter)\s*[, ._/-]*((?:19|20)\d{2})|((?:19|20)\d{2})\s*[, ._/-]*q\s*([1-4]))", text)
    if quarter:
        groups = quarter.groups()
        q, year = (int(groups[0]), int(groups[1])) if groups[0] else (int(groups[2]), int(groups[3])) if groups[2] else (int(groups[5]), int(groups[4]))
        month = 3 * q - 2
        return f"{year:04d}-{month:02d}-01", f"{year:04d}-{month + 2:02d}-{calendar.monthrange(year, month + 2)[1]}", "quarter"
    month_match = re.search(r"(?<!\d)((?:19|20)\d{2})[-/.](0?[1-9]|1[0-2])(?:[-/.]01)?(?!\d)", text)
    if month_match:
        year, month = map(int, month_match.groups())
    else:
        named = re.search(r"(янв\w*|фев\w*|мар\w*|апр\w*|май|мая|июн\w*|июл\w*|авг\w*|сен\w*|окт\w*|ноя\w*|дек\w*)[ ._-]*((?:19|20)?\d{2})(?!\d)", text)
        if named:
            month = _MONTHS[named.group(1)[:3]]
            year = int(named.group(2))
            year += 2000 if year < 100 else 0
        else:
            annual = re.fullmatch(r"(?:(?:отчет\s+за|за|год|период)\s*[:=|]?\s*)?((?:19|20)\d{2})(?:\s*(?:г\.?|год(?:а|у)?))?", text)
            if annual:
                year = int(annual.group(1))
                return f"{year}-01-01", f"{year}-12-31", "year"
            return None
    return f"{year:04d}-{month:02d}-01", f"{year:04d}-{month:02d}-{calendar.monthrange(year, month)[1]}", "month"


def _period_valid(start: str, end: str, granularity: str) -> bool:
    try:
        first, last = date.fromisoformat(start), date.fromisoformat(end)
        if granularity == "year":
            return first.month == first.day == 1 and last == date(first.year, 12, 31)
        if granularity == "month":
            return first.day == 1 and last == date(first.year, first.month, calendar.monthrange(first.year, first.month)[1])
        if granularity == "quarter":
            return first.month in (1, 4, 7, 10) and first.day == 1 and last == date(first.year, first.month + 2, calendar.monthrange(first.year, first.month + 2)[1])
    except (ValueError, TypeError):
        pass
    return False


def _scope(value: object) -> str:
    text = str(value or "").strip()
    return _SCOPE_ALIASES.get(text.upper(), text)


def _unit(value: str) -> str:
    normalized = value.strip().lower()
    if re.fullmatch(r"дн(?:ей|я)?\.?|день|дни", normalized):
        return "дни"
    if re.fullmatch(r"шт\.?|ед\.?|штук", normalized):
        return "шт."
    if re.fullmatch(r"чел\.?|человек", normalized):
        return "чел."
    if normalized == "fte":
        return "FTE"
    if re.fullmatch(r"балл\w*", normalized):
        return "баллы"
    return value.strip()


def _number(value: object, number_format: str = "") -> tuple[float | None, str]:
    if isinstance(value, bool) or value is None:
        return None, ""
    if isinstance(value, (int, float)):
        try:
            number = float(value)
        except (OverflowError, ValueError):
            return None, ""
        return (number * 100, "%") if "%" in number_format else (number, "")
    text = str(value).strip().replace("\u00a0", " ").replace("\u202f", " ").replace("−", "-")
    if text.startswith("="):
        return None, ""
    match = re.fullmatch(r"\s*([+-]?\d[\d ]*(?:[.,]\d+)?)\s*(%|дн(?:ей|я)?\.?|день|дни|шт\.?|чел\.?|fte|балл\w*|(?:тыс\.?\s*)?(?:руб\.?|тенге|тг|₸|₽|usd|kzt|eur))?\s*", text, re.I)
    if not match:
        return None, ""
    numeric = match.group(1).replace(" ", "")
    if re.fullmatch(r"[1-9]\d{0,2}[,.]\d{3}", numeric):
        return None, "неоднозначный разделитель"
    number = float(numeric.replace(",", "."))
    unit = (match.group(2) or "").lower()
    return number, unit


def _finite_number(value: object) -> bool:
    try:
        return not isinstance(value, bool) and isinstance(value, (int, float)) and math.isfinite(value)
    except OverflowError:
        return False


def _literal_value(value: float, quote: str, *, human_grouping=False) -> bool:
    for token in re.findall(r"(?<![\w])[-+]?\d+(?:[ \u00a0\u202f]\d{3})*(?:[.,]\d+)?", quote):
        try:
            observed = float(re.sub(r"[ \u00a0\u202f]", "", token).replace(",", "."))
            if math.isclose(observed, value, rel_tol=1e-9, abs_tol=1e-9):
                return True
            if human_grouping and re.fullmatch(r"[1-9]\d{0,2}[,.]\d{3}", token):
                if math.isclose(float(token.replace(",", "").replace(".", "")), value, rel_tol=1e-9):
                    return True
        except ValueError:
            pass
    return False


def _context_periods(text: str) -> set[tuple[str, str, str]]:
    periods = set()
    for line in text.splitlines():
        if re.search(r"утвержд|протокол|приказ|редакци", line, re.I):
            continue
        parsed = parse_period(line)
        if parsed:
            periods.add(parsed)
        for match in re.finditer(r"(?:за|в|период|отчет\s+за)\s*[:=|]?\s*((?:19|20)\d{2})(?![-/.\d])\s*(?:год\w*|г\.)?", line, re.I):
            parsed = parse_period(match.group(1))
            if parsed:
                periods.add(parsed)
    return periods


def _scope_evidenced(scope: str, text: str, supplied: str) -> bool:
    if supplied and _scope(scope) == _scope(supplied):
        return True
    candidates = [scope] + [alias for alias, canonical in _SCOPE_ALIASES.items() if canonical == scope]
    normalized = " " + normalize_label(text) + " "
    return bool(scope) and any(" " + normalize_label(candidate) + " " in normalized for candidate in candidates)


@dataclass
class _Cell:
    value: object
    row: int
    column: int
    format: str = ""

    @property
    def text(self):
        if self.value is None:
            return ""
        if isinstance(self.value, (int, float)) and not isinstance(self.value, bool) and "%" in self.format:
            return f"{self.value * 100:g}%"
        return str(self.value).strip()


def _unique(items):
    values = list(dict.fromkeys(value for value in items if value))
    return values[0] if len(values) == 1 else None


def _custom(result: IngestionResult, label: str, unit: str = "") -> str:
    code = "custom_" + hashlib.sha256(normalize_label(label).encode()).hexdigest()[:10]
    if not any(item.code == code for item in result.custom_definitions):
        result.custom_definitions.append(MetricDefinition(code, label, unit or "не установлена", "context", [label], confirmed_by_human=False))
    return code


def _ambiguity(result, reason, source, proposed, synthetic):
    identity = hashlib.sha256((source.content_hash + str(source.sheet) + str(source.cell) + source.locator + str(proposed)).encode()).hexdigest()[:18]
    result.ambiguous.append(MetricAmbiguity("ma_" + identity, reason, source, proposed, synthetic))


def _add(result: IngestionResult, *, label: str, value: object, scope: str | None,
         period: tuple | None, source: MetricSource, synthetic: bool, unit: str = "",
         number_format: str = "", extracted_by: str = "rule"):
    number, parsed_unit = _number(value, number_format)
    candidates = resolve_metric(label)
    code = candidates[0] if len(candidates) == 1 else _custom(result, label, unit or parsed_unit) if not candidates else ""
    actual_unit = _unit(unit or parsed_unit)
    proposed = dict(metric=code, label=label, value=number, unit_scope=scope or "", unit=actual_unit,
                    period_start=period[0] if period else "", period_end=period[1] if period else "", granularity=period[2] if period else "")
    reasons = []
    if len(candidates) > 1:
        reasons.append("Неоднозначная метрика: " + ", ".join(candidates))
        proposed["candidates"] = candidates
    elif not candidates:
        reasons.append("Неизвестная метрика: подтвердите название и единицу custom_*")
    if number is None or not math.isfinite(number):
        reasons.append("Число не распознано однозначно; формулы не вычисляются")
    if not scope:
        reasons.append("Не установлено подразделение")
    if not period or not _period_valid(*period):
        reasons.append("Не установлен однозначный полный период")
    definition = METRIC_CATALOG.get(code)
    if definition:
        if actual_unit and code not in {"qa_score", "budget_plan", "budget_fact"} and actual_unit != definition.unit:
            reasons.append("Единица измерения не соответствует выбранной метрике")
        if code in {"budget_plan", "budget_fact"} and not actual_unit:
            reasons.append("Не установлена валюта или единица бюджета")
        if definition.unit == "%" and actual_unit != "%" and "%" not in label and label != code and not re.search(r"доля|процент", label, re.I):
            reasons.append("Не подтверждено, что показатель выражен в процентах")
        if definition.unit != "%" and actual_unit == "%" and code != "qa_score":
            reasons.append("Проценты не соответствуют единице метрики")
        if number is not None and number < 0:
            reasons.append("Отрицательное значение требует проверки")
        if code in {"recs_accepted", "recs_implemented", "assurance_coverage"} and number is not None and number > 100:
            reasons.append("Процент вне диапазона 0–100")
        actual_unit = actual_unit or definition.unit
        proposed["unit"] = actual_unit
    if reasons:
        _ambiguity(result, "; ".join(reasons), source, proposed, synthetic)
        return
    identity = hashlib.sha256((source.content_hash + str(source.sheet) + str(source.cell) + source.locator + code + str(scope) + str(period)).encode()).hexdigest()[:20]
    candidate = MetricValue("mv_" + identity, code, number, scope, *period, source, synthetic, extracted_by, False, actual_unit)
    key = (code, scope, period, synthetic)
    existing = [item for item in result.values if (item.metric, item.unit_scope, (item.period_start, item.period_end, item.granularity), item.is_synthetic) == key]
    conflicts = [item for item in result.ambiguous if item.reason.startswith("Конфликт значений")
                 and (item.proposed.get("metric"), item.proposed.get("unit_scope"),
                      (item.proposed.get("period_start"), item.proposed.get("period_end"), item.proposed.get("granularity")), item.is_synthetic) == key]
    if conflicts:
        _ambiguity(result, "Конфликт значений за один период; выберите авторитетный отчёт", source, proposed, synthetic)
        return
    if any(item.value == number and item.unit == actual_unit for item in existing):
        return
    if existing:
        for item in existing:
            result.values.remove(item)
            _ambiguity(result, "Конфликт значений за один период; выберите авторитетный отчёт", item.source,
                       dict(metric=item.metric, value=item.value, unit_scope=item.unit_scope, period_start=item.period_start, period_end=item.period_end, granularity=item.granularity, unit=item.unit), synthetic)
        _ambiguity(result, "Конфликт значений за один период; выберите авторитетный отчёт", source, proposed, synthetic)
        return
    result.values.append(candidate)


_FIELDS = {
    "metric": {"метрика", "метрики", "показатель", "показатели", "показатель период", "metric", "metric code", "код метрики", "наименование", "наименование показателя"},
    "value": {"значение", "value", "факт", "итого"},
    "period": {"период", "отчетный период", "period", "месяц", "год", "квартал"},
    "scope": {"подразделение", "unit scope", "scope", "отдел"},
    "unit": {"единица", "unit", "единица измерения", "ед измерения", "ед"},
    "period_start": {"period start", "начало периода"},
    "period_end": {"period end", "конец периода"},
    "granularity": {"granularity", "гранулярность", "периодичность"},
}


def _table(result, grid: list[list[_Cell]], *, filename, sheet=None, page=None, table_no=None,
           unit_scope=None, synthetic=False, content_hash="", context=""):
    if not grid:
        return
    context_scope = _unique([_scope(value) for value in re.findall(r"^(?:подразделение|unit_scope)\s*[:=|]\s*([^\n|;]+)$", context, re.I | re.M)])
    inferred_scope = unit_scope or context_scope
    if not inferred_scope and sheet and (sheet.upper() in {"ДНМ", "ДККМ", "ДИТААД", "ДОА", "БВА", "BVA"}):
        inferred_scope = _scope(sheet)
    context_period = _unique([parse_period(line) for line in context.splitlines() if re.search(r"период|отчет\s+за|отчёт\s+за", line, re.I)])
    for header_index, header in enumerate(grid[:30]):
        labels = [normalize_label(cell.text) for cell in header]
        fields = {key: next((index for index, label in enumerate(labels) if label in names), None) for key, names in _FIELDS.items()}
        period_columns = {index: parse_period(cell.value) for index, cell in enumerate(header) if parse_period(cell.value)}
        if fields["metric"] is None and period_columns and header and not header[0].text and 0 not in period_columns:
            fields["metric"] = 0
        metric_columns = {index: cell.text for index, cell in enumerate(header) if resolve_metric(cell.text)}
        long = fields["metric"] is not None and fields["value"] is not None
        wide = fields["metric"] is not None and bool(period_columns)
        transposed = fields["period"] is not None and bool(metric_columns)
        # Unknown labels in a standard period-by-columns table are queued too.
        if fields["period"] is not None and not long and not wide:
            ignored = {value for value in fields.values() if value is not None}
            metric_columns.update({i: cell.text for i, cell in enumerate(header) if i not in ignored and cell.text})
            transposed = bool(metric_columns)
        if not (long or wide or transposed):
            continue
        for row in grid[header_index + 1:]:
            def cell(index):
                return row[index] if index is not None and index < len(row) else _Cell(None, row[0].row if row else 0, 0)
            if not row or not any(item.text for item in row):
                continue
            scope = _scope(cell(fields["scope"]).text) or inferred_scope
            row_unit = cell(fields["unit"]).text
            assignments = []
            if long:
                label = cell(fields["metric"]).text
                if not label or normalize_label(label) in _FIELDS["metric"]:
                    continue
                period = parse_period(cell(fields["period"]).value) if fields["period"] is not None else context_period
                if not period and fields["period_start"] is not None and fields["period_end"] is not None:
                    def iso(value):
                        return value.strftime("%Y-%m-%d") if isinstance(value, (date, datetime)) else str(value or "").strip()
                    start, end = iso(cell(fields["period_start"]).value), iso(cell(fields["period_end"]).value)
                    granularity = cell(fields["granularity"]).text.lower()
                    granularity = {"месяц": "month", "квартал": "quarter", "год": "year"}.get(granularity, granularity)
                    possible = [(start, end, kind) for kind in ("month", "quarter", "year") if _period_valid(start, end, kind)]
                    period = (start, end, granularity) if _period_valid(start, end, granularity) else possible[0] if len(possible) == 1 else None
                assignments = [(label, cell(fields["value"]), period)]
            elif wide:
                label = cell(fields["metric"]).text
                if not label:
                    continue
                assignments = [(label, cell(index), period) for index, period in period_columns.items()]
            else:
                period = parse_period(cell(fields["period"]).value)
                if not period:
                    continue
                assignments = [(label, cell(index), period) for index, label in metric_columns.items()]
            for label, target, period in assignments:
                if target.value in (None, ""):
                    continue
                from openpyxl.utils import get_column_letter
                cell_id = f"{get_column_letter(target.column)}{target.row}" if sheet else f"R{target.row}C{target.column}"
                locator = (f"лист «{sheet}», ячейка {cell_id}" if sheet else f"таблица {table_no}, строка {target.row}, столбец {target.column}")
                if page:
                    locator = f"стр. {page}, " + locator
                locator += f"; заголовок: {header[target.column - 1].text if target.column <= len(header) else ''}"
                quote = " | ".join(item.text for item in row)
                source = MetricSource(filename, quote, sheet, cell_id, page, locator, content_hash)
                _add(result, label=label, value=target.value, scope=scope, period=period,
                     source=source, synthetic=synthetic, unit=row_unit, number_format=target.format)
        return
    result.warnings.append(f"{filename}: таблица {sheet or table_no} не имеет однозначных заголовков метрики/периода; требуется подготовка таблицы или подтверждение человеком.")


def _inspect(filename: str, data: bytes):
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("Отчёт превышает 25 МБ")
    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".docx"}:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Распакованный отчёт превышает 100 МБ")
    tables, fragments, warnings = [], [], []
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            if len(book.worksheets) > 50:
                raise ValueError("В отчёте более 50 листов")
            for sheet in book.worksheets:
                if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 200:
                    raise ValueError("Лист отчёта превышает 10000 строк × 200 столбцов")
                grid = [[_Cell(cell.value, r, c, getattr(cell, "number_format", "")) for c, cell in enumerate(row, 1)] for r, row in enumerate(sheet.iter_rows(), 1)]
                tables.append(dict(grid=grid, sheet=sheet.title, page=None, table_no=None))
        finally:
            book.close()
    elif suffix == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        for number, table in enumerate(doc.tables, 1):
            grid = [[_Cell(cell.text, r, c) for c, cell in enumerate(row.cells, 1)] for r, row in enumerate(table.rows, 1)]
            tables.append(dict(grid=grid, sheet=None, page=None, table_no=number))
        for number, paragraph in enumerate(doc.paragraphs, 1):
            if paragraph.text.strip():
                fragments.append(dict(text=paragraph.text, page=None, locator=f"абзац {number}"))
    elif suffix == ".pdf":
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > 500:
                raise ValueError("PDF-отчёт превышает 500 страниц")
            for number, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    fragments.append(dict(text=text, page=number, locator=f"стр. {number}"))
                else:
                    warnings.append(f"{filename}: стр. {number} без текста; OCR не выполнялся.")
                for table_no, table in enumerate(page.extract_tables(), 1):
                    grid = [[_Cell(value, r, c) for c, value in enumerate(row, 1)] for r, row in enumerate(table, 1)]
                    tables.append(dict(grid=grid, sheet=None, page=number, table_no=table_no))
    else:
        raise ValueError("Поддерживаются отчёты XLSX/XLSM, DOCX и PDF")
    text = "\n".join([item["text"] for item in fragments] + [" | ".join(cell.text for cell in row) for table in tables for row in table["grid"]])
    internal = bool(re.search(r"ТЕСТОВЫЕ ДАННЫЕ|is_synthetic\s*[|:=]\s*(?:true|1)", text, re.I))
    prefix = Path(filename).name.startswith("SYNTH_")
    if prefix != internal:
        raise ValueError(f"{filename}: несогласованные признаки синтетики; нужны SYNTH_ и внутренняя метка ТЕСТОВЫЕ ДАННЫЕ/is_synthetic=true")
    return tables, fragments, warnings, internal


def _consume_text(result, fragments, text_extractor, unit_scope, synthetic):
    if not fragments:
        return
    if text_extractor is None:
        result.warnings.append("Текстовые показатели PDF/Word не анализировались: подключите ИИ-извлечение. Таблицы обработаны правилами.")
        return
    catalog = [asdict(item) for item in METRIC_CATALOG.values()]
    try:
        output = text_extractor(fragments, catalog, unit_scope=unit_scope, is_synthetic=synthetic)
    except Exception:
        result.warnings.append("Не удалось извлечь текстовые показатели через ИИ; таблицы и локальные результаты сохранены.")
        return
    if not isinstance(output, dict):
        result.warnings.append("ИИ-извлечение вернуло неверную структуру.")
        return
    source_map = {item["id"]: item for item in fragments}
    for collection in ("values", "unknown_metrics", "ambiguous"):
        items = output.get(collection, [])
        if not isinstance(items, list):
            continue
        for item in items[:1000]:
            if not isinstance(item, dict) or not isinstance(item.get("source_id"), str) or item["source_id"] not in source_map:
                result.warnings.append("Отклонена цифра ИИ без существующего источника.")
                continue
            fragment = source_map[item["source_id"]]
            quote = item.get("quote")
            if not isinstance(quote, str) or len(quote.strip()) < 8 or quote not in fragment["text"]:
                result.warnings.append("Отклонена цифра ИИ: цитата не совпадает с исходным текстом.")
                continue
            source = MetricSource(fragment["file"], quote, page=fragment.get("page"), locator=fragment["locator"], content_hash=fragment["content_hash"])
            if collection == "ambiguous":
                proposed = item.get("proposed", {})
                _ambiguity(result, str(item.get("reason") or "Требуется подтверждение ИИ-сопоставления"), source, proposed if isinstance(proposed, dict) else {}, synthetic)
                continue
            number = item.get("value")
            if not _finite_number(number) or not _literal_value(float(number), quote):
                result.warnings.append("Отклонена цифра ИИ: значение отсутствует в цитате.")
                continue
            period = (item.get("period_start", ""), item.get("period_end", ""), item.get("granularity", ""))
            label = item.get("metric") if collection == "values" else item.get("name")
            if not isinstance(label, str) or not label:
                _ambiguity(result, "ИИ не определил название метрики", source, item, synthetic)
                continue
            file_text = "\n".join(entry["text"] for entry in fragments if entry["file"] == fragment["file"])
            local_periods = _context_periods(fragment["text"])
            file_periods = _context_periods(file_text)
            supported_periods = local_periods or (file_periods if len(file_periods) == 1 else set())
            scope = _scope(item.get("unit_scope")) or unit_scope
            requested_unit = _unit(str(item.get("unit") or ""))
            reasons = []
            if not _period_valid(*period) or period not in supported_periods:
                reasons.append("Период ИИ не подтверждён текстом отчёта")
            if not _scope_evidenced(scope, file_text, unit_scope):
                reasons.append("Подразделение ИИ не подтверждено текстом или выбором пользователя")
            if requested_unit == "%" and not re.search(r"%|процент|доля", fragment["text"], re.I):
                reasons.append("Процентная единица ИИ отсутствует в исходном фрагменте")
            if reasons:
                codes = resolve_metric(label)
                code = codes[0] if len(codes) == 1 else _custom(result, label, requested_unit) if not codes else ""
                _ambiguity(result, "; ".join(reasons), source,
                           dict(metric=code, value=number, unit_scope=scope, period_start=period[0], period_end=period[1],
                                granularity=period[2], unit=requested_unit, label=label), synthetic)
                continue
            _add(result, label=label, value=number, scope=_scope(item.get("unit_scope")) or unit_scope,
                 period=period if _period_valid(*period) else None, source=source, synthetic=synthetic,
                 unit=str(item.get("unit") or ""), extracted_by="llm")


def ingest_reports(files: list[tuple[str, bytes]], *, unit_scope=None, is_synthetic=False,
                   text_extractor: Callable | None = None) -> IngestionResult:
    result = IngestionResult()
    inspected = [(filename, data, _inspect(filename, data)) for filename, data in files]
    namespaces = {entry[2][3] for entry in inspected}
    if len(namespaces) > 1 or (namespaces and next(iter(namespaces)) != is_synthetic):
        raise ValueError("Нельзя смешивать реальные и синтетические отчёты или менять их происхождение флагом.")
    fragments = []
    for filename, data, (tables, raw_fragments, warnings, synthetic) in inspected:
        fingerprint = hashlib.sha256(data).hexdigest()
        result.warnings.extend(warnings)
        context = "\n".join(item["text"] for item in raw_fragments[:30])
        for table in tables:
            if table["sheet"] and table["sheet"].lower() in {"_metadata", "metadata", "метаданные"}:
                continue
            header_context = "\n".join(" | ".join(cell.text for cell in row) for row in table["grid"][:8])
            table_context = "\n".join(item["text"] for item in raw_fragments if item.get("page") == table["page"]) if table["page"] else context
            _table(result, **table, filename=filename, unit_scope=_scope(unit_scope), synthetic=synthetic,
                   content_hash=fingerprint, context=table_context + "\n" + header_context)
        for index, fragment in enumerate(raw_fragments, 1):
            fragments.append({**fragment, "id": f"mr_{fingerprint[:16]}_{index}", "file": filename, "content_hash": fingerprint})
    _consume_text(result, fragments, text_extractor, _scope(unit_scope), is_synthetic)
    return result


def confirm_ambiguity(result: IngestionResult, ambiguity_id: str, *, metric: str, unit_scope: str,
                      period_start: str, period_end: str, granularity: str, value=None, unit=None,
                      custom_name=None) -> MetricValue:
    item = next((item for item in result.ambiguous if item.id == ambiguity_id), None)
    if item is None:
        raise ValueError("Неоднозначность уже разрешена или не найдена")
    value = item.proposed.get("value") if value is None else value
    if not _finite_number(value) or not _literal_value(value, item.source.quote, human_grouping=True):
        raise ValueError("Подтверждаемое число должно присутствовать в исходной цитате")
    if not unit_scope.strip() or not _period_valid(period_start, period_end, granularity):
        raise ValueError("Укажите подразделение и полный месяц/квартал/год")
    definition = METRIC_CATALOG.get(metric) or next((entry for entry in result.custom_definitions if entry.code == metric), None)
    if definition is None:
        raise ValueError("Выберите метрику из каталога или предложенный custom_*")
    if metric.startswith("custom_"):
        if not (custom_name or definition.name) or not unit:
            raise ValueError("Для custom_* нужны описание и единица измерения")
        definition.name = custom_name or definition.name
        definition.unit = unit
        definition.confirmed_by_human = True
    candidate = MetricValue(item.id.replace("ma_", "mv_"), metric, float(value), _scope(unit_scope),
                            period_start, period_end, granularity, item.source, item.is_synthetic,
                            "human", True, unit or definition.unit)
    result.values.append(candidate)
    if item.reason.startswith("Конфликт значений"):
        result.ambiguous[:] = [entry for entry in result.ambiguous if not (entry.reason.startswith("Конфликт значений")
                              and entry.proposed.get("metric") == item.proposed.get("metric")
                              and entry.proposed.get("unit_scope") == item.proposed.get("unit_scope")
                              and entry.proposed.get("period_start") == item.proposed.get("period_start")
                              and entry.proposed.get("period_end") == item.proposed.get("period_end"))]
    else:
        result.ambiguous.remove(item)
    return candidate
