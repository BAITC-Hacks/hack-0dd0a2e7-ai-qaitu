from __future__ import annotations

import hashlib
import io
import re
import zipfile
from pathlib import Path
from typing import BinaryIO, Literal

from .models import Document, Fragment


Period = Literal["before", "after"]
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_PAGES = 1000
MAX_FRAGMENTS = 20_000
MAX_TEXT_CHARS = 5_000_000
MAX_SHEETS = 50
MAX_ROWS = 10_000
MAX_COLUMNS = 200

# Numbers inside a cross-reference or date must not start a new clause.
_CLAUSE_MARKER = re.compile(r"(?<!\S)(\d{1,2}(?:\.\d{1,3})*)[.)](?=\s*[А-Яа-яЁёA-Za-z])")
_CLAUSE_START = re.compile(r"^(\d{1,2}(?:\.\d{1,3})*)[.)](?:\s*|$)")
_SUBITEM = re.compile(r"^([а-яёa-z])[.)]\s*", re.I)
_REFERENCE_END = re.compile(r"(?:\bп|\bпп|\bст|\bразд)\.$", re.I)


def _clean(text: object) -> str:
    return re.sub(r"\s+", " ", "" if text is None else str(text)).strip()


def _pdf_page_text(page: object) -> str:
    """Recover reading lines when a PDF's text layer emits one word per line."""
    plain = page.extract_text() or ""
    lines = [line.strip() for line in plain.splitlines() if line.strip()]
    if len(lines) > 20 and sum(len(line.split()) <= 1 for line in lines) / len(lines) > 0.65:
        layout = page.extract_text(extraction_mode="layout") or ""
        if layout.strip():
            return layout
    return plain


def _pdf_fragments(text: str) -> list[str]:
    return re.split(r"\n\s*\n|\n(?=\s*(?:\d+(?:\.\d+)+\.|[а-я]\.)\s+)", text, flags=re.I)


def _fragment(
    name: str, period: Period, locator: str, text: str, n: int, fingerprint: str
) -> Fragment:
    # Include both content and name: two different files sharing a basename
    # cannot overwrite each other's citations. Repeated identical uploads dedupe.
    name_key = hashlib.sha256(name.encode("utf-8")).hexdigest()[:10]
    return Fragment(f"{period}:{name_key}:{fingerprint[:20]}:{n}", name, _clean(text), locator, period)


def _split_clauses(text: str) -> list[tuple[str, int, int]]:
    """Return exact text spans; offsets are 1-based in the original paragraph."""
    starts = [0]
    for match in _CLAUSE_MARKER.finditer(text):
        if not match.start():
            continue
        prefix = text[:match.start()].rstrip()
        # Inline clauses usually follow a sentence/semicolon. Newlines are
        # sufficient too. Do not split “согласно п. 5.3.2. настоящего ...”.
        gap = text[len(prefix):match.start()]
        if (prefix[-1:] in ".;:!?" or "\n" in gap) and not _REFERENCE_END.search(prefix):
            starts.append(match.start())
    return [
        (text[start:end], start + 1, end)
        for start, end in zip(starts, starts[1:] + [len(text)])
        if _clean(text[start:end])
    ]


class _Collector:
    def __init__(self, name: str, period: Period, fingerprint: str):
        self.name, self.period, self.fingerprint = name, period, fingerprint
        self.fragments: list[Fragment] = []
        self.clause = ""
        self.characters = 0

    def add(self, text: str, locator: str) -> None:
        spans = _split_clauses(text)
        for part, start, end in spans:
            cleaned = _clean(part)
            explicit = _CLAUSE_START.match(cleaned)
            subitem = _SUBITEM.match(cleaned)
            if explicit:
                self.clause = explicit.group(1)
            reference = self.clause
            if subitem and self.clause:
                reference += subitem.group(1).lower()
            position = locator
            if len(spans) > 1:
                position += f", симв. {start}–{end}"
            if reference:
                position += f", п. {reference}"
            self.characters += len(cleaned)
            if len(self.fragments) >= MAX_FRAGMENTS or self.characters > MAX_TEXT_CHARS:
                raise ValueError("Документ превышает лимит текста прототипа; разделите его на части.")
            self.fragments.append(_fragment(
                self.name, self.period, position, cleaned,
                len(self.fragments) + 1, self.fingerprint,
            ))


def _check_archive(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 10_000 or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
            raise ValueError("Распакованный документ превышает безопасный лимит 100 МБ.")


def _metadata(fragments: list[Fragment], fingerprint: str) -> dict[str, str]:
    cover = []
    for fragment in fragments[:30]:
        if _CLAUSE_START.match(fragment.text):
            break
        cover.append(fragment.text)
    text = " ".join(cover).replace("«", "").replace("»", "")
    result = {"sha256": fingerprint}
    for key, pattern in {
        "edition": r"редакци[яи]\s*(?:(?:№|No|N)\s*)?(\d+)",
        "protocol": r"протокол[а-я]*\s*(?:от\s+[^№]{0,60})?(?:№|No|N)\s*(\d+)",
        "date": r"(\d{1,2}\s+[а-яё]+\s+20\d{2})(?:\s*г\b|\s*года\b)",
    }.items():
        match = re.search(pattern, text, re.I)
        if match:
            result[key] = match.group(1)
    return result


def extract_document(file: BinaryIO, name: str, period: Period) -> Document:
    """Extract ordered, traceable text locally; unsupported coverage is explicit."""
    suffix = Path(name).suffix.lower()
    if suffix not in {".pdf", ".docx", ".xlsx", ".xlsm"}:
        raise ValueError(f"Формат {suffix or 'без расширения'} не поддерживается")
    if period not in {"before", "after"}:
        raise ValueError("Период должен быть before или after")
    data = file.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("Файл превышает лимит 25 МБ.")
    fingerprint = hashlib.sha256(data).hexdigest()
    collector = _Collector(name, period, fingerprint)
    warnings: list[str] = []
    if suffix != ".pdf":
        _check_archive(data)

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("PDF защищён паролем. Загрузите доступную для чтения копию.")
        if len(reader.pages) > MAX_PAGES:
            raise ValueError(f"PDF превышает лимит {MAX_PAGES} страниц.")
        empty_pages = []
        for page_no, page in enumerate(reader.pages, 1):
            page_text = _pdf_page_text(page)
            if not _clean(page_text):
                empty_pages.append(page_no)
                continue
            paragraphs = _pdf_fragments(page_text)
            for block, paragraph in enumerate(paragraphs, 1):
                collector.add(paragraph, f"стр. {page_no}, блок {block}")
        if empty_pages:
            pages = ", ".join(map(str, empty_pages[:20]))
            more = "…" if len(empty_pages) > 20 else ""
            warnings.append(f"Страницы без извлекаемого текста: {pages}{more}. OCR не выполнялся; покрытие неполное.")

    elif suffix == ".docx":
        from docx import Document as DocxDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = DocxDocument(io.BytesIO(data))
        paragraph_no = table_no = 0
        auto_numbering = False
        # doc.paragraphs followed by doc.tables changes ownership context. Walk
        # body children instead so a table stays between surrounding headings.
        for element in doc.element.body.iterchildren():
            if element.tag.endswith("}p"):
                paragraph_no += 1
                paragraph = Paragraph(element, doc)
                style = paragraph.style
                numbered = bool(element.xpath("./w:pPr/w:numPr"))
                visited = set()
                while style is not None and style.style_id not in visited:
                    visited.add(style.style_id)
                    numbered |= bool(style.element.xpath("./w:pPr/w:numPr"))
                    style = style.base_style
                auto_numbering |= numbered
                collector.add(paragraph.text, f"абзац {paragraph_no}")
            elif element.tag.endswith("}tbl"):
                table_no += 1
                table = Table(element, doc)
                for row_no, row in enumerate(table.rows, 1):
                    cells = []
                    previous_cell = None
                    for cell in row.cells:
                        # Merged cells repeat the same XML cell in python-docx.
                        if cell._tc is previous_cell:
                            continue
                        previous_cell = cell._tc
                        cells.append(_clean(cell.text))
                    text = " | ".join(cells)
                    if any(cells):
                        collector.add(text, f"таблица {table_no}, строка {row_no}")
        if auto_numbering:
            warnings.append("Автонумерация Word не восстановлена: используйте номера абзацев и цитаты; номера пунктов могут отсутствовать.")
        if doc.element.body.xpath(".//w:ins | .//w:del | .//w:txbxContent"):
            warnings.append("В Word есть исправления или текстовые блоки: их полное извлечение не поддерживается. Примите исправления и проверьте текст.")

    else:
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if len(book.worksheets) > MAX_SHEETS:
                raise ValueError(f"Книга превышает лимит {MAX_SHEETS} листов.")
            for sheet in book.worksheets:
                collector.clause = ""
                if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                    raise ValueError(f"Лист «{sheet.title}» превышает лимит {MAX_ROWS} строк × {MAX_COLUMNS} столбцов.")
                for row_no, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if row_no > MAX_ROWS or len(row) > MAX_COLUMNS:
                        raise ValueError(f"Лист «{sheet.title}» превышает лимиты прототипа.")
                    cells = [_clean(value) for value in row]
                    if any(cells):
                        collector.add(" | ".join(cells), f"лист «{sheet.title}», строка {row_no}")
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if any(b"<f" in archive.read(path) for path in archive.namelist() if path.startswith("xl/worksheets/") and path.endswith(".xml")):
                    warnings.append("Формулы Excel не вычисляются: извлечены сохранённые значения; отсутствующий кэш формул даёт пустые ячейки.")
        finally:
            book.close()

    if not collector.fragments:
        detail = " ".join(warnings)
        raise ValueError(f"В документе «{name}» не найден извлекаемый текст. {detail}".strip())
    return Document(name=name, period=period, fragments=collector.fragments,
                    warnings=warnings, metadata=_metadata(collector.fragments, fingerprint))


def document_from_lines(name: str, period: Period, lines: list[str]) -> Document:
    fingerprint = hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()
    collector = _Collector(name, period, fingerprint)
    for i, line in enumerate(lines, 1):
        collector.add(line, f"строка {i}")
    return Document(name=name, period=period, fragments=collector.fragments,
                    metadata=_metadata(collector.fragments, fingerprint))
