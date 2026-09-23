"""Offline analytical workbook and shared executive-report content.

Document text is data: Excel strings never become formulas or external links.
Exports consume the full AnalysisResult, independently of screen filters.
"""
from __future__ import annotations

from collections import Counter
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from .confidence import LEVELS, confidence_counts, confidence_level, metric_lines
from .models import AnalysisResult
from .reporting import AI_COMPARISON_LABELS, COVERAGE_LABELS, FINDING_LABELS, NORM_LABELS, ROLE_LABELS, STATUS_LABELS, ai_review_overview, collect_sources


DISCLAIMER = ("Выводы — кандидаты для экспертной проверки. Отсутствие назначения в загруженных документах "
              "не доказывает прекращение функции; несколько владельцев не доказывают дублирование. "
              "Уверенность — эвристическая оценка, не вероятность нарушения и не тяжесть последствий.")
SHEET_NAMES = ("Сводка", "Изменения структуры", "Риски", "ИИ-сравнение", "Матрица функций", "Потери", "Переданные функции", "Новые функции", "Источники")
CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def clean_text(value) -> str:
    return CONTROL_CHARS.sub("\ufffd", str(value)).replace("\r\n", "\n").replace("\r", "\n")


def clip(value: str, size: int) -> str:
    value = clean_text(value)
    return value if len(value) <= size else value[:size].rstrip() + "… [сокращено]"


def level_label(score: float) -> str:
    return LEVELS[confidence_level(score)].split(" ", 1)[1]


def score_of(item) -> float | None:
    if item.assessment:
        return item.assessment.score
    score = getattr(item, "confidence", None)
    return min(.99, max(0, score)) if score is not None else None


def document_lines(result: AnalysisResult, document_packs: dict | None = None) -> list[str]:
    lines = []
    for period, label in (("before", "До"), ("after", "После")):
        records = (document_packs or {}).get(period, [])
        if not records:
            names = dict.fromkeys(s.document for s in collect_sources(result) if s.period == period)
            records = [{"name": name} for name in names]
        for record in records:
            metadata = record.get("metadata", {})
            details = ", ".join(f"{title} {metadata[key]}" for key, title in (
                ("edition", "редакция №"), ("protocol", "протокол №"), ("date", "дата")) if metadata.get(key))
            lines.append(f"{label}: {record['name']}" + (f" ({details})" if details else ""))
    return lines


def overview(result: AnalysisResult) -> list[tuple[str, str | int]]:
    units = Counter(item.status for item in result.unit_changes)
    rows = Counter(item.status for item in result.matrix_rows)
    risks = Counter(item.kind for item in result.findings)
    sources = collect_sources(result)
    return [
        ("Документов до / после", f"{result.coverage.get('documents_before', len({s.document for s in sources if s.period == 'before'}))} / "
                                     f"{result.coverage.get('documents_after', len({s.document for s in sources if s.period == 'after'}))}"),
        ("Строк функциональной матрицы", len(result.matrix_rows)),
        ("Структура: сохранено / преобразовано", f"{units['preserved']} / {units['transformed']}"),
        ("Структура: появилось / не найдено в перечне", f"{units['created']} / {units['removed']}"),
        ("Функции: сохранено / передано", f"{rows['preserved']} / {rows['moved']}"),
        ("Функции: изменена формулировка / новое закрепление", f"{rows['changed']} / {rows['new']}"),
        ("Функции: закрепление не найдено / недостаточно данных", f"{rows['lost']} / {rows['unknown']}"),
        ("Кандидаты на потерю", risks['loss']),
        ("Потенциальные пересечения", risks['duplicate']),
        ("Потенциальные конфликты ролей", risks['conflict']),
    ]


def executive_conclusion(result: AnalysisResult) -> str:
    rows = Counter(row.status for row in result.matrix_rows)
    losses = sum(f.kind == "loss" for f in result.findings)
    return (f"Сохранено назначений: {rows['preserved']}; передано или изменён набор владельцев: {rows['moved']}; "
            f"изменена формулировка: {rows['changed']}. Кандидатов на потерю: {losses}. "
            "Индикаторы требуют проверки по исходным пунктам и границам ответственности."
            if result.matrix_rows else "Недостаточно извлечённых назначений для содержательного сравнения. Проверьте документы и предупреждения.")


def excel_report(result: AnalysisResult, *, document_packs: dict | None = None, is_demo: bool = False) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "QAITU — анализ организационных изменений"
    workbook.properties.creator = "QAITU"
    sources = collect_sources(result)
    source_rows: dict[str, int] = {}

    def add_sheet(name, headers, rows, *, percentages=(), source_column=None):
        sheet = workbook.create_sheet(name)
        sheet.append([*headers, "Часть записи"])
        for logical in rows:
            # Avoid Excel's silent 32767-character truncation. 15000 code points
            # also fit if every character uses two UTF-16 code units.
            values = [clean_text(v) if isinstance(v, str) else v for v in logical]
            parts = [(len(v) + 14999) // 15000 if isinstance(v, str) else 1 for v in values]
            for part in range(max(parts, default=1)):
                output = [v[part * 15000:(part + 1) * 15000] if isinstance(v, str) and len(v) > 15000 else v for v in values]
                output.append(f"{part + 1}/{max(parts)}")
                sheet.append(output)
                for cell, value in zip(sheet[sheet.max_row], output):
                    if isinstance(value, str):
                        cell.data_type = "s"  # Includes =, +, -, @ and error-looking text.
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cell.font = Font(name="Calibri", size=11)
                for column in percentages:
                    sheet.cell(sheet.max_row, column).number_format = "0%"
                if name == "Источники":
                    source_rows.setdefault(logical[0], sheet.max_row)
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17364A")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 34
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for i, header in enumerate(headers, 1):
            width = 62 if any(word in header for word in ("Текст", "Функция", "Обоснование", "Рекомендация", "Ограничения", "Значение", "Метрики")) else 30
            sheet.column_dimensions[get_column_letter(i)].width = 16 if header == "Уверенность" else width
        sheet.column_dimensions[get_column_letter(len(headers) + 1)].width = 15
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
        sheet.oddFooter.center.text = "QAITU · &P / &N"
        if source_column:
            for row in sheet.iter_rows(min_row=2):
                cell = row[source_column - 1]
                target = source_rows.get(cell.value)
                if target:
                    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'Источники'!A{target}")
                    cell.font = Font(name="Calibri", size=11, color="156680", underline="single")
        return sheet

    # Build catalog first for real internal hyperlinks; reorder to reader order below.
    add_sheet("Источники", ["Источник ID", "Период", "Документ", "Локатор", "Текст источника"],
              [(s.id, "До" if s.period == "before" else "После", s.document, s.locator, s.text) for s in sources])
    summary = [("Режим", "Синтетический контрольный пример" if is_demo else "Загруженные документы"), *overview(result),
               ("Вывод", executive_conclusion(result)), ("Как читать выводы", DISCLAIMER),
               ("Матрица", "Одна строка на назначение и владельца; связанные назначения обеих редакций имеют одинаковый ID строки матрицы. Фильтруйте по периоду, подразделению, статусу и числовой уверенности."),
               ("Полнота экспорта", "Все находки, назначения и источники, независимо от фильтров экрана. Пустая уверенность означает, что оценка не вычислена. Очень длинные ячейки разделены на последовательные части записи; управляющие символы заменены символом замены. Для длинных текстов используйте автоподбор высоты строки."),
               ("Шкала", "90–99% очень высокая; 70–89% высокая; 40–69% требует проверки; менее 40% слабый сигнал. Уверенность не равна ущербу."),
               ("Полнота комплекта после", "Заявлена пользователем" if result.analysis_context.get("after_complete_user_declared") else "Не подтверждена")]
    summary.extend(ai_review_overview(result))
    summary.extend(("Документ", line) for line in document_lines(result, document_packs))
    summary.extend(("Ограничение обработки", warning) for warning in result.warnings)
    summary.extend((COVERAGE_LABELS.get(key, key), value) for key, value in result.coverage.items())
    for kind, bands in confidence_counts(result.findings).items():
        summary.extend((f"{FINDING_LABELS[kind]} · {LEVELS[band].split(' ', 1)[1]}", bands[band]) for band in LEVELS)
    add_sheet("Сводка", ["Показатель", "Значение"], summary)
    add_sheet("Изменения структуры", ["Статус", "До", "После", "Источники ID"], [
        (STATUS_LABELS[item.status], item.before or "—", item.after or "—", "\n".join(s.id for s in item.sources)) for item in result.unit_changes])
    row_map = {row.id: row for row in result.matrix_rows}
    risks = []
    for i, finding in enumerate(result.findings, 1):
        value = finding.assessment
        row = row_map.get(finding.matrix_row_id)
        risks.append((f"R{i}", FINDING_LABELS[finding.kind], finding.title, score_of(finding), level_label(score_of(finding)),
                      finding.matrix_row_id, "\n".join(dict.fromkeys(f.unit for f in row.before + row.after)) if row else "См. источники",
                      finding.explanation, finding.recommendation,
                      "\n".join(value.reasons) if value else "", "\n".join(value.limitations) if value else "",
                      "\n".join(metric_lines(value)) if value else "", "\n".join(s.id for s in finding.sources),
                      "\n".join(s.id for s in value.evidence) if value else "", value.method if value else ""))
    add_sheet("Риски", ["Риск ID", "Тип", "Название", "Уверенность", "Уровень уверенности", "ID строки матрицы", "Владельцы",
                       "Обоснование", "Рекомендация", "Причины оценки", "Ограничения", "Метрики оценки", "Источники ID", "Источники сравнения ID", "Метод"], risks, percentages=(4,))

    source_map = {source.id: source for source in sources}
    ai_rows = []
    for index, comparison in enumerate(result.ai_review.get("comparisons", []), 1):
        for evidence in comparison.get("evidence", []) or [{}]:
            source_id = evidence.get("source_id", "")
            source = source_map.get(source_id)
            ai_rows.append((f"AI{index}", AI_COMPARISON_LABELS.get(comparison.get("kind"), "Сопоставление"),
                            comparison.get("title", ""), comparison.get("explanation", ""), comparison.get("recommendation", ""),
                            "\n".join(comparison.get("before_source_ids", [])), "\n".join(comparison.get("after_source_ids", [])),
                            source_id, evidence.get("quote", ""), "До" if source and source.period == "before" else "После" if source else "Не найден в каталоге",
                            source.document if source else "", source.locator if source else ""))
    add_sheet("ИИ-сравнение", ["ИИ ID", "Вид", "Название", "Обоснование", "Рекомендация", "До · источники ID", "После · источники ID",
                              "Источник ID", "Точная цитата", "Период", "Документ", "Локатор"], ai_rows, source_column=8)

    headers = ["ID строки матрицы", "Статус", "Функция", "Уверенность", "Тип нормы", "Период", "Подразделение", "Текст назначения",
               "Роль", "Область", "Владелец известен", "Источник ID", "Документ", "Локатор", "Контекст ID", "Примечания"]

    def assignments(rows):
        for row in rows:
            for period, functions in (("До", row.before), ("После", row.after)):
                for f in functions:
                    yield (row.id, STATUS_LABELS.get(row.status, row.status), row.label, score_of(row), NORM_LABELS[f.norm_type], period,
                           f.unit, f.text, ROLE_LABELS.get(f.role, f.role), f.scope, "Да" if f.owner_known else "Нет", f.source.id,
                           f.source.document, f.source.locator, "\n".join(s.id for s in f.context_sources), "\n".join(row.notes))

    for name, rows in (("Матрица функций", result.matrix_rows), ("Потери", [row for row in result.matrix_rows if row.status == "lost"]),
                       ("Переданные функции", [row for row in result.matrix_rows if row.status == "moved"]),
                       ("Новые функции", [row for row in result.matrix_rows if row.status == "new"])):
        add_sheet(name, headers, assignments(rows), percentages=(4,), source_column=12)
    for index, name in enumerate(SHEET_NAMES):
        workbook.move_sheet(name, offset=index - workbook.sheetnames.index(name))
    workbook.active = 0
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
