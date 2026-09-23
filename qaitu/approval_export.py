"""Offline approval sheet with all decisions, source quotes and audit history."""
from __future__ import annotations

from io import BytesIO
from xml.sax.saxutils import escape

from .exports import clean_text
from .models import AnalysisResult
from .review import STATUSES, current_decisions, finding_id, gate_status, package_id, review_items


def approval_sections(result: AnalysisResult, history: list[dict]):
    gate = gate_status(result, history)
    yield 0, "QAITU · Лист согласования"
    yield 1, gate["label"]
    yield 2, f"Комплект: {package_id(result)}"
    yield 2, f"Вопросов: {gate['total']} · открыто: {gate['open']} · закрыто: {gate['closed']}"
    yield 2, gate["limitation"]
    yield 2, "Имена введены пользователями; это не проверка личности и не электронная подпись. Исправления отмечаются вручную; этот файл не изменяет исходное положение."
    for reason in gate["reasons"]:
        yield 2, reason
    current = current_decisions(history)
    for finding in review_items(result):
        key = finding_id(finding)
        decision = current.get(key, {})
        yield 1, f"{key} · {finding.title}"
        yield 2, finding.explanation
        yield 2, "Решение: " + STATUSES[decision.get("status", "found")]
        if decision:
            yield 2, f"Кто: {decision['actor']} · когда (UTC): {decision['timestamp']} · ответственный: {decision['assignee'] or 'не назначен'}"
            yield 2, "Комментарий: " + decision["comment"]
            yield 2, f"Экспертный приоритет: {decision['impact']} (не балл матрицы рисков)"
            if decision.get("correction_reference"):
                yield 2, "Место исправления: " + decision["correction_reference"]
        for source in finding.sources:
            yield 2, f"Источник {source.id} · {source.document} · {source.locator} · {'до' if source.period == 'before' else 'после'}"
            yield 2, source.text
        previous = [event for event in history if event["finding_id"] == key]
        if previous:
            yield 2, "История решений (UTC):"
            for event in previous:
                yield 2, f"#{event['revision']} · {event['timestamp']} · {event['actor']} · {STATUSES[event['status']]} · {event['comment']}"


def approval_docx(result: AnalysisResult, history: list[dict]) -> bytes:
    from docx import Document
    from docx.shared import Pt
    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)
    for level, text in approval_sections(result, history):
        text = clean_text(text)
        if level < 2:
            document.add_heading(text, level=level)
        else:
            document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def approval_pdf(result: AnalysisResult, history: list[dict]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate
    from .pdf_report import PDF_BUILD_LOCK, _register_fonts
    with PDF_BUILD_LOCK:
        _register_fonts()
        styles = [ParagraphStyle("title", fontName="QaituBold", fontSize=18, leading=24, spaceAfter=12),
                  ParagraphStyle("heading", fontName="QaituBold", fontSize=11, leading=15, spaceBefore=12, spaceAfter=7, keepWithNext=True),
                  ParagraphStyle("body", fontName="Qaitu", fontSize=9, leading=12, spaceAfter=6, splitLongWords=True)]
        output = BytesIO()
        story = [Paragraph(escape(clean_text(text)).replace("\n", "<br/>"), styles[level])
                 for level, text in approval_sections(result, history)]

        def footer(canvas, doc):
            canvas.setFont("Qaitu", 8)
            canvas.drawString(18 * mm, 12 * mm, f"QAITU · Лист согласования · {doc.page}")

        SimpleDocTemplate(output, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                          topMargin=18 * mm, bottomMargin=20 * mm,
                          title="QAITU · Лист согласования", author="QAITU").build(story, onFirstPage=footer, onLaterPages=footer)
        return output.getvalue()
