import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, Mock, patch

from docx import Document
from openpyxl import Workbook, load_workbook

from qaitu.metrics_catalog import METRIC_CATALOG, resolve_metric
from qaitu.metrics_ingest import confirm_ambiguity, ingest_reports, parse_period
from scripts.gen_synthetic_reports import generate_reports


def xlsx(rows, sheet="ДККМ", formats=None, synthetic=False):
    book = Workbook()
    page = book.active
    page.title = sheet
    for row in rows:
        page.append(row)
    for cell, number_format in (formats or {}).items():
        page[cell].number_format = number_format
    if synthetic:
        metadata = book.create_sheet("_metadata")
        metadata.append(["is_synthetic", True])
        metadata.append(["label", "ТЕСТОВЫЕ ДАННЫЕ"])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def word(paragraphs=(), table=None):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        output = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                output.cell(r, c).text = str(value)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


class MetricsExtractionTests(unittest.TestCase):
    def test_catalog_includes_all_nineteen_distinct_codes(self):
        self.assertEqual(len(METRIC_CATALOG), 19)
        self.assertEqual(METRIC_CATALOG["findings_total"].better, "context")
        self.assertEqual(METRIC_CATALOG["audits_done"].normalize_by, "fte_actual")
        self.assertEqual(set(resolve_metric("Количество проверок")), {"audits_done", "audits_planned"})

    def test_periods_year_quarter_and_russian_month(self):
        cases = {
            "2022": ("2022-01-01", "2022-12-31", "year"),
            "за 2023 год": ("2023-01-01", "2023-12-31", "year"),
            "Q1 2023": ("2023-01-01", "2023-03-31", "quarter"),
            "1 кв. 2023": ("2023-01-01", "2023-03-31", "quarter"),
            "янв.23": ("2023-01-01", "2023-01-31", "month"),
            "февраль 2024": ("2024-02-01", "2024-02-29", "month"),
            "2023-12": ("2023-12-01", "2023-12-31", "month"),
        }
        for label, expected in cases.items():
            with self.subTest(label=label):
                self.assertEqual(parse_period(label), expected)
        for label in ("утверждено 15.03.2023", "2023-02-15", "2023-02-01", "2022 vs 2023", "приказ за 2023 год", "Обсуждались результаты в компании 2023"):
            self.assertIsNone(parse_period(label), label)

    def test_synthetic_extraction_exceeds_95_percent_and_every_value_is_sourced(self):
        with tempfile.TemporaryDirectory() as directory:
            manifest = generate_reports(directory)
            payloads = [(Path(path).name, Path(path).read_bytes()) for path in manifest["files"]]
            result = ingest_reports(payloads, is_synthetic=True)
            self.assertGreaterEqual(len(result.values), .95 * manifest["expected_values"])
            self.assertEqual(len(result.values), 2052)
            self.assertFalse(result.ambiguous)
            self.assertFalse(result.warnings)
            workbooks = {name: load_workbook(io.BytesIO(data), data_only=True) for name, data in payloads}
            for value in result.values:
                self.assertTrue(value.is_synthetic)
                self.assertTrue(value.source.file.startswith("SYNTH_"))
                self.assertTrue(value.source.content_hash)
                self.assertTrue(value.source.quote)
                cell_value = workbooks[value.source.file][value.source.sheet][value.source.cell].value
                self.assertAlmostEqual(value.value, cell_value)
                self.assertIn(value.metric, value.source.quote)
            for book in workbooks.values():
                book.close()
            self.assertEqual(manifest["events"][1]["units"], ["БВА:ИТ-аудит"])
            self.assertTrue(manifest["controls"]["W1"]["verified"])

    def test_flexible_wide_quarter_table_preserves_exact_cells(self):
        data = xlsx([["Показатель", "1 кв. 2023", "Q2 2023"], ["Задержка отчетов", 0, 7]])
        result = ingest_reports([("report.xlsx", data)])
        self.assertEqual([item.value for item in result.values], [0, 7])
        self.assertEqual([item.source.cell for item in result.values], ["B2", "C2"])
        self.assertEqual(result.values[0].period_end, "2023-03-31")
        self.assertEqual(result.values[0].unit_scope, "ДККМ")

    def test_transposed_metrics_and_percentage_format_are_supported(self):
        data = xlsx([["Период", "audits_done", "plan_completion"], ["янв.23", 5, .8]], formats={"C2": "0%"})
        result = ingest_reports([("report.xlsx", data)])
        self.assertEqual({item.metric: item.value for item in result.values}, {"audits_done": 5, "plan_completion": 80})
        self.assertIn("80%", result.values[1].source.quote)
        self.assertEqual(result.values[1].unit, "%")

    def test_explicit_start_and_end_dates_define_a_complete_period(self):
        data = xlsx([["metric", "value", "period_start", "period_end", "granularity", "unit_scope"],
                     ["audits_done", 7, "2023-01-01", "2023-03-31", "quarter", "ДНМ"]])
        result = ingest_reports([("report.xlsx", data)])
        self.assertEqual(len(result.values), 1)
        self.assertEqual(result.values[0].granularity, "quarter")
        self.assertEqual(result.values[0].period_end, "2023-03-31")

    def test_unknown_metric_and_missing_owner_are_queued_without_guessing(self):
        data = xlsx([["Показатель", "Период", "Значение"], ["Среднее время согласования", "2023", 12]], sheet="Свод")
        result = ingest_reports([("report.xlsx", data)])
        self.assertEqual(result.values, [])
        self.assertEqual(len(result.ambiguous), 1)
        self.assertEqual(len(result.custom_definitions), 1)
        self.assertFalse(result.custom_definitions[0].confirmed_by_human)
        self.assertIn("подразделение", result.ambiguous[0].reason)
        code = result.custom_definitions[0].code
        confirmed = confirm_ambiguity(result, result.ambiguous[0].id, metric=code, unit_scope="ДККМ",
                                      period_start="2023-01-01", period_end="2023-12-31", granularity="year", unit="часы", custom_name="Время согласования")
        self.assertTrue(confirmed.confirmed_by_human)
        self.assertNotIn(code, METRIC_CATALOG)
        self.assertEqual(len(result.ambiguous), 0)

    def test_ambiguous_alias_requires_human_choice_and_cannot_invent_number(self):
        data = xlsx([["Показатель", "2023"], ["Количество проверок", 8]])
        result = ingest_reports([("report.xlsx", data)])
        item = result.ambiguous[0]
        with self.assertRaisesRegex(ValueError, "исходной цитате"):
            confirm_ambiguity(result, item.id, metric="audits_done", value=999, unit_scope="ДККМ", period_start="2023-01-01", period_end="2023-12-31", granularity="year")
        value = confirm_ambiguity(result, item.id, metric="audits_done", value=8, unit_scope="ДККМ", period_start="2023-01-01", period_end="2023-12-31", granularity="year")
        self.assertEqual(value.value, 8)

    def test_bare_accepted_count_cannot_be_silently_treated_as_percent(self):
        data = xlsx([["Показатель", "2023"], ["Принято рекомендаций", 8]])
        result = ingest_reports([("report.xlsx", data)])
        self.assertEqual(result.values, [])
        self.assertIn("процент", result.ambiguous[0].reason)

    def test_ambiguous_thousands_separator_requires_explicit_confirmation(self):
        data = xlsx([["Показатель", "2023"], ["Проверок проведено", "1,234"]])
        result = ingest_reports([("report.xlsx", data)])
        self.assertFalse(result.values)
        item = result.ambiguous[0]
        confirmed = confirm_ambiguity(result, item.id, metric="audits_done", value=1234, unit_scope="ДККМ", period_start="2023-01-01", period_end="2023-12-31", granularity="year")
        self.assertEqual(confirmed.value, 1234)
        self.assertTrue(confirmed.confirmed_by_human)

    def test_conflicting_reports_remain_ambiguous_even_after_third_value(self):
        files = [(f"report{i}.xlsx", xlsx([["Показатель", "2023"], ["Проверок проведено", value]])) for i, value in enumerate((5, 8, 9))]
        result = ingest_reports(files)
        self.assertEqual(result.values, [])
        self.assertEqual(len(result.ambiguous), 3)
        item = result.ambiguous[1]
        confirm_ambiguity(result, item.id, metric="audits_done", value=8, unit_scope="ДККМ", period_start="2023-01-01", period_end="2023-12-31", granularity="year")
        self.assertEqual([item.value for item in result.values], [8])
        self.assertFalse(result.ambiguous)

    def test_word_table_has_row_and_cell_provenance(self):
        data = word(["Подразделение: ДНМ"], [["Показатель", "2023"], ["Проверок проведено", "5"]])
        result = ingest_reports([("report.docx", data)])
        self.assertEqual(len(result.values), 1)
        value = result.values[0]
        self.assertEqual(value.unit_scope, "ДНМ")
        self.assertEqual(value.source.cell, "R2C2")
        self.assertIn("таблица 1", value.source.locator)
        self.assertEqual(value.source.quote, "Проверок проведено | 5")

    def test_pdf_table_has_page_and_cell_provenance(self):
        page = SimpleNamespace(extract_text=lambda: "Подразделение: ДНМ", extract_tables=lambda: [[["Показатель", "2023"], ["Проверок проведено", "5"]]])
        pdf = MagicMock()
        pdf.__enter__.return_value.pages = [page]
        with patch("pdfplumber.open", return_value=pdf):
            result = ingest_reports([("report.pdf", b"pdf-test")])
        self.assertEqual(result.values[0].source.page, 1)
        self.assertEqual(result.values[0].source.cell, "R2C2")
        self.assertIn("стр. 1", result.values[0].source.locator)

    def test_text_callback_requires_literal_number_period_and_owner(self):
        text = "ДККМ. За 2023 год проведено 8 проверок."
        data = word([text])
        def extractor(fragments, catalog, **kwargs):
            return {"values": [dict(metric="audits_done", value=8, unit_scope="ДККМ", period_start="2023-01-01", period_end="2023-12-31", granularity="year", source_id=fragments[0]["id"], quote=text, unit="шт.")]}
        result = ingest_reports([("report.docx", data)], text_extractor=extractor)
        self.assertEqual(len(result.values), 1)
        self.assertEqual(result.values[0].extracted_by, "llm")
        self.assertEqual(result.values[0].source.quote, text)
        def invented(fragments, catalog, **kwargs):
            output = extractor(fragments, catalog, **kwargs)
            output["values"][0].update(unit_scope="Несуществующий отдел", period_start="2024-01-01", period_end="2024-12-31")
            return output
        rejected = ingest_reports([("report.docx", data)], text_extractor=invented)
        self.assertEqual(rejected.values, [])
        self.assertIn("Период ИИ", rejected.ambiguous[0].reason)
        self.assertIn("Подразделение ИИ", rejected.ambiguous[0].reason)

    def test_text_callback_cannot_invent_or_rewrite_number(self):
        text = "ДККМ. За 2023 год проведено 8 проверок."
        def extractor(fragments, catalog, **kwargs):
            return {"values": [dict(metric="audits_done", value=42, unit_scope="ДККМ", period_start="2023-01-01", period_end="2023-12-31", granularity="year", source_id=fragments[0]["id"], quote=text, unit="шт.")]}
        result = ingest_reports([("report.docx", word([text]))], text_extractor=extractor)
        self.assertFalse(result.values)
        self.assertTrue(any("значение отсутствует" in warning for warning in result.warnings))

    def test_text_callback_cannot_add_percentage_unit_to_an_absolute_count(self):
        text = "ДККМ. За 2023 год принято 8 рекомендаций."
        def extractor(fragments, catalog, **kwargs):
            return {"values": [dict(metric="recs_accepted", value=8, unit_scope="ДККМ", period_start="2023-01-01", period_end="2023-12-31", granularity="year", source_id=fragments[0]["id"], quote=text, unit="%")]}
        result = ingest_reports([("report.docx", word([text]))], text_extractor=extractor)
        self.assertFalse(result.values)
        self.assertIn("Процентная единица", result.ambiguous[0].reason)

    def test_mixed_or_relabelled_namespaces_are_rejected(self):
        rows = [["Показатель", "2023"], ["Проверок проведено", 5]]
        real = ("report.xlsx", xlsx(rows))
        synth = ("SYNTH_report.xlsx", xlsx(rows, synthetic=True))
        for files, flag in [([real, synth], False), ([real], True), ([synth], False), ([("SYNTH_report.xlsx", real[1])], True)]:
            with self.subTest(flag=flag), self.assertRaises(ValueError):
                ingest_reports(files, is_synthetic=flag)


if __name__ == "__main__":
    unittest.main()
