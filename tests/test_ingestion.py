import io
import unittest
from unittest.mock import patch

from docx import Document as WordDocument
from openpyxl import Workbook

from qaitu.extractors import document_from_lines, extract_document


class IngestionTests(unittest.TestCase):
    @staticmethod
    def word(document):
        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)
        return stream

    def test_same_filename_different_content_has_distinct_sources(self):
        first = document_from_lines("Положение.docx", "before", ["Первый текст"])
        second = document_from_lines("Положение.docx", "before", ["Второй текст"])
        self.assertNotEqual(first.fragments[0].id, second.fragments[0].id)
        self.assertEqual(first.fragments[0].id, document_from_lines("Положение.docx", "before", ["Первый текст"]).fragments[0].id)

    def test_docx_preserves_table_between_paragraphs(self):
        document = WordDocument()
        document.add_paragraph("Отдел закупок")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Функция"
        table.cell(0, 1).text = "Проводит закупки"
        document.add_paragraph("Отдел аудита")
        result = extract_document(self.word(document), "rules.docx", "after")
        self.assertEqual([fragment.text for fragment in result.fragments], ["Отдел закупок", "Функция | Проводит закупки", "Отдел аудита"])
        self.assertIn("таблица 1, строка 1", result.fragments[1].locator)
        self.assertIn("абзац 2", result.fragments[2].locator)

    def test_embedded_clauses_are_split_with_precise_source_spans(self):
        document = WordDocument()
        text = "3.9. Рабочие места находятся в филиалах. 3.10.Работники подчиняются директору. 3.11. Сотрудники ведут отчеты."
        document.add_paragraph(text)
        result = extract_document(self.word(document), "rules.docx", "after")
        self.assertEqual(len(result.fragments), 3)
        self.assertEqual(result.fragments[1].text, "3.10.Работники подчиняются директору.")
        self.assertIn("п. 3.10", result.fragments[1].locator)
        self.assertIn(f"симв. {text.index('3.10.') + 1}–", result.fragments[1].locator)
        self.assertEqual(len({fragment.id for fragment in result.fragments}), 3)

    def test_subitems_inherit_clause_without_rewriting_quote(self):
        result = document_from_lines("rules.docx", "after", ["5.3.3. Взаимодействует в части:", "а. использования результатов проверки;", "б. выявления рисков."])
        self.assertEqual(result.fragments[1].text, "а. использования результатов проверки;")
        self.assertIn("п. 5.3.3а", result.fragments[1].locator)
        self.assertIn("п. 5.3.3б", result.fragments[2].locator)

    def test_cross_reference_does_not_split_paragraph(self):
        result = document_from_lines("rules.docx", "after", ["5.3.3. Исполняет работу согласно п. 5.3.2. настоящего Положения."])
        self.assertEqual(len(result.fragments), 1)

    def test_metadata_uses_cover_instead_of_dates_referenced_in_body(self):
        result = document_from_lines("rules.docx", "after", [
            "УТВЕРЖДЕНО", "Протокол No 7", "от «23» декабря 2022 года",
            "ПОЛОЖЕНИЕ (редакция No9)", "1. Общие положения",
            "1.1. На основании акта от 22 марта 2012 года.",
        ])
        self.assertEqual(result.metadata["edition"], "9")
        self.assertEqual(result.metadata["protocol"], "7")
        self.assertEqual(result.metadata["date"], "23 декабря 2022")
        body_only = document_from_lines("rules.docx", "after", ["1. Общие положения", "1.1. На основании акта от 22 марта 2012 года."])
        self.assertNotIn("date", body_only.metadata)

    def test_word_auto_numbering_is_explicitly_unsupported(self):
        document = WordDocument()
        document.add_paragraph("Проводит аудит", style="List Number")
        result = extract_document(self.word(document), "rules.docx", "after")
        self.assertTrue(any("Автонумерация" in warning for warning in result.warnings))

    def test_spreadsheet_preserves_zero_false_and_column_positions(self):
        book = Workbook()
        book.active.append(["Функция", None, 0, False])
        book.active.append(["После", None, 1])
        stream = io.BytesIO()
        book.save(stream)
        stream.seek(0)
        result = extract_document(stream, "matrix.xlsx", "after")
        self.assertEqual(result.fragments[0].text, "Функция | | 0 | False")
        self.assertIn("строка 2", result.fragments[1].locator)

    def test_xlsx_and_xlsm_are_both_accepted(self):
        book = Workbook()
        book.active.append(["Департамент аудита", "проверяет закупки"])
        stream = io.BytesIO()
        book.save(stream)
        data = stream.getvalue()
        for extension in ("xlsx", "xlsm"):
            with self.subTest(extension=extension):
                result = extract_document(io.BytesIO(data), f"matrix.{extension}", "after")
                self.assertEqual(result.fragments[0].text, "Департамент аудита | проверяет закупки")
                self.assertIn("лист", result.fragments[0].locator)

    def test_spreadsheet_limits_fail_without_silent_truncation(self):
        book = Workbook()
        book.active.cell(10_001, 1, "За пределами лимита")
        stream = io.BytesIO()
        book.save(stream)
        stream.seek(0)
        with self.assertRaisesRegex(ValueError, "превышает лимит"):
            extract_document(stream, "matrix.xlsx", "after")

    def test_scanned_pdf_page_produces_coverage_warning(self):
        from types import SimpleNamespace
        pages = [SimpleNamespace(extract_text=lambda: "Отдел аудита"), SimpleNamespace(extract_text=lambda: "")]
        with patch("pypdf.PdfReader", return_value=SimpleNamespace(pages=pages, is_encrypted=False)):
            result = extract_document(io.BytesIO(b"pdf"), "rules.pdf", "after")
        self.assertTrue(any("OCR" in warning and "2" in warning for warning in result.warnings))


if __name__ == "__main__":
    unittest.main()
