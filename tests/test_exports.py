import json
import unittest
from copy import deepcopy
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader

from qaitu.analyzer import analyze_documents
from qaitu.confidence import assessment
from qaitu.demo import demo_documents
from qaitu.exports import SHEET_NAMES, excel_report
from qaitu.models import AnalysisResult, Finding, Fragment, Function, MatrixRow
from qaitu.pdf_report import pdf_report
from qaitu.reporting import collect_sources


def pdf(data):
    reader = PdfReader(BytesIO(data))
    return reader, "\n".join(page.extract_text() for page in reader.pages)


class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        before, after = demo_documents()
        cls.result = analyze_documents(before, after)

    def test_short_pdf_has_cyrillic_summary_structure_findings_and_no_catalog(self):
        reader, text = pdf(pdf_report(self.result, is_demo=True))
        self.assertTrue(5 <= len(reader.pages) <= 10)
        self.assertIn("Анализ реорганизации", reader.pages[0].extract_text())
        self.assertIn("Общий вывод", reader.pages[0].extract_text())
        self.assertIn("Изменения структуры", reader.pages[1].extract_text())
        self.assertIn("Синтетический контрольный пример", text)
        self.assertIn("Уверенность:", text)
        self.assertIn("Передачи ответственности", text)
        self.assertIn("Новые закрепления функций", text)
        self.assertNotIn("Приложение А", text)
        self.assertNotIn("Полный каталог источников", text)
        self.assertGreaterEqual(len(reader.outline), 5)
        fonts = reader.pages[0]["/Resources"]["/Font"].get_object()
        self.assertTrue(any("/FontFile2" in item.get_object().get("/FontDescriptor", {}) for item in fonts.values()))

    def test_full_pdf_keeps_all_sources_and_links_to_catalog(self):
        reader, text = pdf(pdf_report(self.result, full=True))
        self.assertIn("Приложение А. Полная функциональная матрица", text)
        self.assertIn("Приложение Б. Полный каталог источников", " ".join(text.split()))
        self.assertIn("Конец полного каталога источников.", text)
        self.assertTrue(any(page.mediabox.width > page.mediabox.height for page in reader.pages))
        for page in reader.pages:
            if page.mediabox.width > page.mediabox.height:
                self.assertLessEqual(page.extract_text().count("Статус / ID"), 1)
        self.assertLess(reader.pages[-1].mediabox.width, reader.pages[-1].mediabox.height)
        for source in collect_sources(self.result):
            self.assertIn(source.id, text)
            self.assertIn(source.locator, text)
        links = [item.get_object() for page in reader.pages for item in page.get("/Annots", [])]
        self.assertTrue(any("/Dest" in item for item in links))
        self.assertFalse(any(item.get("/A", {}).get("/S") == "/URI" for item in links))

    def test_short_pdf_selects_top_ten_and_excludes_weak_but_full_keeps_them(self):
        result = deepcopy(self.result)
        template = result.findings[0]
        result.findings = [replace(template, title=f"SIGNAL-{i:02d}", confidence=.80 + i / 100,
                                   assessment=assessment(.80 + i / 100)) for i in range(12)]
        result.findings.append(replace(template, title="WEAK-ONLY-SENTINEL", confidence=.1, assessment=assessment(.1)))
        _, short = pdf(pdf_report(result))
        self.assertIn("Показано 10 из 13", short)
        self.assertIn("SIGNAL-11", short)
        self.assertNotIn("SIGNAL-00", short)
        self.assertNotIn("WEAK-ONLY-SENTINEL", short)
        _, full = pdf(pdf_report(result, full=True))
        self.assertIn("WEAK-ONLY-SENTINEL", full)

    def test_excel_eight_sheets_filters_freeze_numeric_scores_and_internal_links(self):
        workbook = load_workbook(BytesIO(excel_report(self.result)))
        self.assertEqual(workbook.sheetnames, list(SHEET_NAMES))
        for sheet in workbook:
            self.assertTrue(sheet.freeze_panes)
            self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
        risks = workbook["Риски"]
        self.assertEqual(risks.max_row - 1, len(self.result.findings))
        self.assertIsInstance(risks["D2"].value, float)
        self.assertEqual(risks["D2"].number_format, "0%")
        matrix = workbook["Матрица функций"]
        self.assertEqual(matrix.max_row - 1, sum(len(r.before) + len(r.after) for r in self.result.matrix_rows))
        self.assertEqual(matrix["G1"].value, "Подразделение")
        self.assertEqual({row[0].value for row in matrix.iter_rows(min_row=2)}, {r.id for r in self.result.matrix_rows})
        for row in matrix.iter_rows(min_row=2):
            self.assertTrue(row[11].hyperlink.location.startswith("'Источники'!A"))
        sources = workbook["Источники"]
        self.assertEqual(sources.max_row - 1, len(collect_sources(self.result)))
        self.assertEqual({row[0].value: row[4].value for row in sources.iter_rows(min_row=2)},
                         {s.id: s.text for s in collect_sources(self.result)})
        for name, status in (("Потери", "lost"), ("Переданные функции", "moved"), ("Новые функции", "new")):
            self.assertEqual(workbook[name].max_row - 1, sum(len(r.before) + len(r.after) for r in self.result.matrix_rows if r.status == status))

    def test_excel_hostile_cells_are_text_not_formulas(self):
        payloads = ['=HYPERLINK("https://example.invalid","click")', '+SUM(1,1)', '-1+1', '@SUM(1,1)', '#REF!', '\t=cmd', '<b>текст</b>\x00']
        result = AnalysisResult([], [], [], sources=[Fragment(f"id-{i}", value, value, value, "before") for i, value in enumerate(payloads)])
        workbook = load_workbook(BytesIO(excel_report(result)))
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")
                    self.assertFalse(cell.hyperlink)
        self.assertEqual(workbook["Источники"]["E2"].value, payloads[0])
        self.assertIn("\ufffd", workbook["Источники"]["E8"].value)

    def test_excel_very_long_source_is_split_without_silent_truncation(self):
        text = "Текст с кириллицей " * 2400 + "КОНЕЦ-ИСТОЧНИКА"
        source = Fragment("long", "Документ", text, "п. 1", "before")
        workbook = load_workbook(BytesIO(excel_report(AnalysisResult([], [], [], sources=[source]))))
        rows = list(workbook["Источники"].iter_rows(min_row=2))
        self.assertGreater(len(rows), 1)
        self.assertEqual("".join(row[4].value or "" for row in rows), text)
        self.assertEqual(rows[-1][5].value, f"{len(rows)}/{len(rows)}")

    def test_pdf_long_table_row_and_xml_like_text_never_disappear(self):
        text = '<link href="https://example.invalid">не ссылка</link> & <b>не разметка</b> ' + "Длинное назначение функции. " * 380 + "КОНЕЦ-ДЛИННОГО-ПУНКТА"
        source = Fragment("long", "Документ", text, "п. 8", "before")
        f = Function("f1", "ДККМ", text, source)
        row = MatrixRow("long-row", text, [f], [], "lost")
        result = AnalysisResult([], [], [], matrix_rows=[row])
        reader, extracted = pdf(pdf_report(result, full=True))
        self.assertIn("КОНЕЦ-ДЛИННОГО-ПУНКТА", extracted)
        self.assertIn('<link href="https://example.invalid">', extracted)
        self.assertIn("Конец полного каталога источников", extracted)
        self.assertGreater(len(reader.pages), 7)

    def test_empty_results_export_and_no_confidence_is_not_zero(self):
        result = AnalysisResult([], [], [])
        for full in (True, False):
            _, text = pdf(pdf_report(result, full=full))
            self.assertIn("Недостаточно извлечённых назначений", text)
        workbook = load_workbook(BytesIO(excel_report(result)))
        self.assertEqual(workbook["Риски"].max_row, 1)
        self.assertEqual(workbook["Матрица функций"].max_row, 1)
        result = deepcopy(self.result)
        result.matrix_rows[0].assessment = None
        workbook = load_workbook(BytesIO(excel_report(result)))
        self.assertIsNone(workbook["Матрица функций"]["D2"].value)

    def test_exports_preserve_result_and_document_metadata(self):
        original = json.dumps(self.result.to_dict(), sort_keys=True, ensure_ascii=False)
        packs = {"before": [{"name": "old.pdf", "metadata": {"edition": "8"}}],
                 "after": [{"name": "new.docx", "metadata": {"edition": "9"}}]}
        _, text = pdf(pdf_report(self.result, document_packs=packs))
        self.assertIn("редакция № 8", text)
        self.assertIn("редакция № 9", text)
        excel_report(self.result, document_packs=packs)
        self.assertEqual(original, json.dumps(self.result.to_dict(), sort_keys=True, ensure_ascii=False))

    def test_concurrent_downloads_do_not_mix_document_content(self):
        def build(marker):
            result = AnalysisResult([], [], [], sources=[Fragment(marker, marker + ".docx", marker, "п. 1", "before")])
            return pdf(pdf_report(result, full=True))[1]
        with ThreadPoolExecutor(max_workers=2) as pool:
            first, second = list(pool.map(build, ("USER-ALPHA-SENTINEL", "USER-BETA-SENTINEL")))
        self.assertIn("USER-ALPHA-SENTINEL", first)
        self.assertNotIn("USER-BETA-SENTINEL", first)
        self.assertIn("USER-BETA-SENTINEL", second)
        self.assertNotIn("USER-ALPHA-SENTINEL", second)


if __name__ == "__main__":
    unittest.main()
